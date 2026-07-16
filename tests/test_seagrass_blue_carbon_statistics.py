import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPT_PATH = REPO_ROOT / "Seagrass_blue_carbon_statistics_v1.0.py"


class SeagrassBlueCarbonStatisticsTest(unittest.TestCase):
    def create_workbook(self, path: Path, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "observations"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)

    def test_runs_with_documented_workbook_structure(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.xlsx"
            output_path = temp_path / "output.xlsx"
            self.create_workbook(
                input_path,
                [
                    "site",
                    "station",
                    "transect",
                    "sample_id",
                    "species",
                    "biomass_g_m2",
                    "carbon_fraction",
                    "area_ha",
                ],
                [
                    ["Sumber Kima", "ST01", "T1", "SK-01", "Enhalus acoroides", 825, 0.36, 1.5],
                    ["Sumber Kima", "ST01", "T1", "SK-02", "Thalassia hemprichii", 640, 0.34, 1.5],
                    ["Banyuwedang", "ST02", "T2", "BW-01", "Cymodocea rotundata", 455, 0.31, 2.0],
                ],
            )

            result = subprocess.run(
                ["python", str(SCRIPT_PATH), str(input_path), "--output", str(output_path)],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Processed 3 observations.", result.stdout)
            self.assertTrue(output_path.exists())

            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "overall_summary",
                    "site_summary",
                    "species_summary",
                    "observations_enriched",
                ],
            )
            enriched = workbook["observations_enriched"]
            headers = [cell.value for cell in enriched[1]]
            carbon_stock_index = headers.index("carbon_stock_mg_c_ha") + 1
            total_carbon_index = headers.index("total_carbon_mg_c") + 1
            # carbon_stock = biomass_g_m2 * carbon_fraction * 0.01; total_carbon = stock * area_ha
            expected_values = {
                2: (2.97, 4.455),
                3: (2.176, 3.264),
                4: (1.4105, 2.821),
            }
            for row_number, (expected_stock, expected_total) in expected_values.items():
                self.assertAlmostEqual(
                    enriched.cell(row=row_number, column=carbon_stock_index).value,
                    expected_stock,
                )
                self.assertAlmostEqual(
                    enriched.cell(row=row_number, column=total_carbon_index).value,
                    expected_total,
                )

    def test_fails_when_required_columns_are_missing(self):
        scenarios = [
            (["site", "species", "biomass_g_m2"], [["Sumber Kima", "Enhalus acoroides", 825]]),
            (["site"], [["Sumber Kima"]]),
            (["notes"], [["incomplete workbook"]]),
        ]
        for index, (headers, rows) in enumerate(scenarios, start=1):
            with self.subTest(headers=headers):
                with tempfile.TemporaryDirectory() as temp_dir:
                    temp_path = Path(temp_dir)
                    input_path = temp_path / f"invalid-{index}.xlsx"
                    self.create_workbook(input_path, headers, rows)

                    result = subprocess.run(
                        ["python", str(SCRIPT_PATH), str(input_path)],
                        capture_output=True,
                        text=True,
                        check=False,
                    )

                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn("missing required columns", result.stderr)

    def test_fails_when_biomass_is_negative(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "negative-biomass.xlsx"
            self.create_workbook(
                input_path,
                ["site", "species", "biomass_g_m2", "carbon_fraction"],
                [["Sumber Kima", "Enhalus acoroides", -1, 0.36]],
            )

            result = subprocess.run(
                ["python", str(SCRIPT_PATH), str(input_path)],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("negative biomass_g_m2", result.stderr)


if __name__ == "__main__":
    unittest.main()
